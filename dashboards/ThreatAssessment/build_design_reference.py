from pathlib import Path
import os,json,zipfile,tempfile
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
root=Path(r'\\cscso1fsappv01\home\jaskew\HTOC\dashboards\ThreatAssessment')
work=Path(tempfile.mkdtemp(prefix='htoc_design_reference_'));p=Path(r'Z:\HTOC\Data_Analytics\Data\Threat Assessment Scores\ThreatAssessI_W\Dashboard\Dashobard.pbix')
with zipfile.ZipFile(p) as z:
 theme=json.loads(z.read('Report/StaticResources/RegisteredResources/HTOC_IW_Dark.json').decode('utf-8-sig'));layout=json.loads(z.read('Report/Layout').decode('utf-16-le'))
theme['name']='HTOC I&W Dashboard';theme_path=work/'HTOC I&W Dashboard Theme.json';theme_path.write_text(json.dumps(theme,indent=2),encoding='utf-8')
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7);sec.header_distance=sec.footer_distance=Inches(.25)
for name,size in [('Normal',11),('Title',23),('Heading 1',16),('Heading 2',12)]:
 s=doc.styles[name];s.font.name='Segoe UI';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0);s.font.underline=False
 if name!='Normal':s.font.bold=True
 s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=1.05
 if name.startswith('Heading'):s.paragraph_format.space_before=Pt(10);s.paragraph_format.keep_with_next=True
 for r in s.element.xpath('.//w:color'):
  for attr in ['themeColor','themeTint','themeShade']:
   r.attrib.pop(qn('w:'+attr),None)
 for x in s.element.xpath('.//w:pBdr'):x.getparent().remove(x)
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('HTOC I&W Dashboard  |  Design reference  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor(0,0,0)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
def para(text,boldlead=None):
 p=doc.add_paragraph()
 if boldlead and text.startswith(boldlead):p.add_run(boldlead).bold=True;p.add_run(text[len(boldlead):])
 else:p.add_run(text)
 return p
def heading(t,level=1):return doc.add_paragraph(t,'Heading '+str(level))
def table(headers,rows,widths,center_cols=()):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for i,label in enumerate(headers):t.rows[0].cells[i].text=label
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for i,val in enumerate(row):cells[i].text=str(val)
 for ri,row in enumerate(t.rows):
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=cell._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   pr.append(borders);m=OxmlElement('w:tcMar')
   for edge,v in [('top',65),('bottom',65),('left',90),('right',90)]:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(v));el.set(qn('w:type'),'dxa');m.append(el)
   pr.append(m);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'303366' if ri==0 else ('F5F5F8' if ri%2==0 else 'FFFFFF'));pr.append(sh)
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0;p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ci in center_cols else WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:r.font.name='Segoe UI';r.font.size=Pt(9.5);r.font.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 return t
# Page 1
p0=doc.add_paragraph('HTOC Dashboard Style and Feature Guide','Title')
para('Use this reference and the accompanying Power BI theme to reproduce the HTOC I&W Dashboard. It records the visual standards, layout, interactions, and refresh limits that future designs should preserve.')
p0=para('Baseline captured 26 October 2026  |  Version 1.0  |  Source Dashobard.pbix');p0.runs[0].font.size=Pt(9)
heading('Colors')
colors=[('Canvas','#0D0D0D','Page background'),('Panels and hover cards','#171722','Cards, containers, table backgrounds, and definition popups'),('Bars and partner blocks','#303366','PRISM bars and partner blocks'),('Text and data labels','#D9D9D9','Titles, values, labels, footnotes, and information icons'),('Panel outlines','#2A3A4A','Subtle borders with 8 pixel rounded corners where used')]
t=table(['Role','Color','Use'],colors,[1.75,1.05,4.3],center_cols=(1,))
for i,row in enumerate(colors,1):
 c=t.rows[i].cells[1];sh=OxmlElement('w:shd');sh.set(qn('w:fill'),row[1][1:]);c._tc.get_or_add_tcPr().append(sh)
 for r in c.paragraphs[0].runs:r.font.color.rgb=RGBColor.from_string('000000' if row[1]=='#D9D9D9' else 'FFFFFF')
heading('Typography')
fonts=[('Dashboard title','Segoe UI','24 pt','HTOC regular; I&W Dashboard bold'),('Panel and chart titles','Segoe UI Semibold','13 pt','Use consistent title case'),('KPI values','Segoe UI Semibold','28 pt','One clear value per card'),('Date value','Segoe UI Semibold','22 pt','dd MMM yyyy'),('Table headings','Segoe UI Semibold','10 pt','Centered'),('Table data','Segoe UI','10 pt','Centered; keep each record on one row'),('PRISM value labels','Segoe UI Semibold','11 pt','Visible values; consistent number formatting'),('Chart footnotes and update time','Segoe UI','8 pt','Compact and readable; update time on one line'),('Column information glyph','Segoe UI Symbol','9 pt','Circled information icon beside the heading'),('Hover card headings','Segoe UI','15 pt','Match existing definition cards'),('Hover card body','Segoe UI','12 pt','Plain language for nontechnical readers'),('Assessment note and tag tables','Segoe UI','11 pt','Show source text for the hovered Indicator')]
table(['Element','Typeface','Size','Rule'],fonts,[1.8,1.75,.7,2.85])
para('The theme JSON provides a reusable baseline. Direct visual formatting in the PBIX supplies several sizes and component settings, so importing the theme alone does not reproduce the full dashboard.')
# Page 2
doc.add_page_break();heading('Layout and interactions')
para('Use one visible page at 1120 by 680 pixels with Fit to page. Keep the page background dark and preserve the balanced panel layout. Hidden pages are reserved for hover cards.')
by={json.loads(v['config'])['name']:v for v in layout['sections'][0]['visualContainers']}
components=[('Title','dashboard_title'),('Indicator count','kpi_0'),('Highest PRISM score','kpi_3'),('Date','overview_asof'),('Last updated','last_updated_timestamp'),('Reported Indicators panel','reported_panel'),('Indicator Details panel','indicator_details_frame'),('Partner panel','severity_panel'),('PRISM panel','prism_panel')]
geometry=[]
for label,name in components:
 v=by[name];geometry.append((label,round(v['x']),round(v['y']),round(v['width']),round(v['height'])))
table(['Component','X','Y','Width','Height'],geometry,[3.1,1,1,1,1],center_cols=(1,2,3,4))
heading('Indicator Details',2)
para('Use one native table with one vertical scrollbar. Columns are Indicator, Severity, PRISM Score, Indicator Type, Threat Actor, Last Observed, and OpDiv. Center headings and data. Sorting must keep every field for an Indicator together. Do not split the table into separate column visuals.')
heading('Partner blocks',2)
para('Use equal-area partner blocks that shrink and rearrange to fit the available panel without a scrollbar or pagination. Block size does not represent Indicator count. Split the source Partners list into distinct partner-to-Indicator associations so an Indicator can match several partners.')
para('Selecting a partner filters Indicator Details and the PRISM score chart; other KPI cards and the reported list retain page-wide values. For one Indicator, keep its partner blocks visible and the same record and score displayed.')
heading('Hover information',2)
para('Assessment notes and source tags appear on row hover. One native table shares the row tooltip across its data cells. Each column gets a themed definition card through the information icon beside its title. Keep icons clear of text and preserve sorting and shared scrolling.')
para('Use a small native card with a hidden information-glyph measure and a manually assigned tooltip page for each column. This provides themed hover cards.')
heading('Reported Indicators and chart notes',2)
para('The reported panel lists distinct Indicators from I&W_Yes sheets in the five-calendar-day file window, including today. Keep its footnote and the PRISM explanation inside their respective panels.')
# Page 3
doc.add_page_break();heading('Definitions and refresh behavior')
definitions=[('Indicator','The IP address, domain, URL, file, or other observable being assessed for potential threats.'),('Severity','The priority assigned in the PRISM threat assessment, such as Critical, High, Medium, or Low.'),('PRISM Score','PRISM is an internal threat assessment model that evaluates all verified evidence about an Indicator and assigns a score from 0 to 1,000, with 1,000 representing the highest assessed maliciousness and 0 the lowest.'),('Indicator Type','The kind of Indicator, such as an IP address, domain, URL, or file hash.'),('Threat Actor','The actor or group linked to the Indicator in the source assessment. A blank value means the source does not list an actor.'),('Last Observed','When the source assessment says the Indicator was most recently seen.'),('OpDiv','The operating division or divisions listed for the Indicator.'),('Date','The data snapshot date, displayed as dd MMM yyyy.'),('Last updated','The most recent data-load timestamp from Snapshot[Refreshed At], displayed as Last updated: dd MMM yyyy HH:mm. A layout edit does not change it.')]
table(['Field','Definition'],definitions,[1.25,5.85])
heading('Current data and refresh behavior',2)
para('The source folder is ThreatAssessI_W. Daily files use ThreatAssessI_W_YYYYMMDD.xlsx. The active dashboard is Dashboard\\Dashobard.pbix. Current Indicator data comes from I&W_No; reported history comes from I&W_Yes. The dashboard remains a local import report.')
para('Automatic file pickup and scheduled refresh are not enabled. The current Indicator and Snapshot partitions contain inline imported data; a new file drop does not update the open dashboard. The reported-history folder query evaluates its five-day window when refreshed. Future live deployment must implement and validate ingestion and refresh separately. Keep Task Scheduler disabled unless the user requests a change.')
heading('Rules for future changes',2)
for text in [
'Preserve the current original and all user edits outside the requested change. Keep the HTOC Big Board reference PBIX read-only.',
'Derive dates, counts, scores, and partners from data. Preserve the distinction between data date and last updated time; do not substitute the current clock time.',
'Test one and multiple Indicators, shared sorting and scrolling, per-column hover definitions, row assessment notes and tags, and growth from six to ten partner blocks.',
'Check full labels, consistent colors, readable text, and unwanted text-box scrollbars after changing fonts or component sizes.',
'Keep backups and obsolete test dashboards in Deprecated folders. Open the final PBIX in the installed Power BI version after package or model changes.'
]:
 p0=doc.add_paragraph(text,'List Bullet');p0.paragraph_format.space_after=Pt(4)
para('Reuse the accompanying HTOC I&W Dashboard Theme.json for colors and base fonts. Recreate the documented layout, measures, relationships, and hover behavior separately; these features are not installed by a theme file.')
path=work/'HTOC Dashboard Style and Feature Guide.docx';doc.save(path)
(root/'design_reference_state.json').write_text(json.dumps({'work':str(work),'docx':str(path),'theme':str(theme_path)},indent=2))
print(json.dumps({'work':str(work),'docx':str(path),'theme':str(theme_path)}))
