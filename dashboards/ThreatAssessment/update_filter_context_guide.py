from pathlib import Path
import json,hashlib,copy
from docx import Document
from docx.shared import Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Row
root=Path(r'\\cscso1fsappv01\home\jaskew\HTOC\dashboards\ThreatAssessment');context_state_path=Path(r'C:\Users\jaskew\AppData\Local\Temp\htoc_filter_context_ku2zekev\state.json');context=json.loads(context_state_path.read_text());work=Path(context['work']);source=Path(r'Z:\HTOC\Data_Analytics\Data\Threat Assessment Scores\ThreatAssessI_W\Dashboard\Design Reference\HTOC Dashboard Style and Feature Guide.docx');original=source.read_bytes();doc=Document(source)
for p in doc.paragraphs:
 for r in p.runs:r.text=r.text.replace('Version 1.0','Version 1.1')
 if p.text.startswith('Selecting a partner filters Indicator Details and the PRISM score chart;'):
  p.add_run(' The header shows Showing X of Y Indicators and the selected partner. Clear filter restores all Indicators without resetting table sorting.')
t=doc.tables[1];template=next(r for r in t.rows if r.cells[0].text=='Table data');new_tr=copy.deepcopy(template._tr);template._tr.addnext(new_tr);row=_Row(new_tr,t)
for cell,value in zip(row.cells,['Filter context and reset','Segoe UI','10 / 9 pt','Context uses 10 pt; Clear filter uses 9 pt Semibold.']):
 cell.text=value
 for p in cell.paragraphs:
  p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
  for r in p.runs:r.font.name='Segoe UI';r.font.size=Pt(9.5);r.font.color.rgb=RGBColor(0,0,0);r.bold=False
for i,row in enumerate(t.rows):
 if i:
  for cell in row.cells:
   pr=cell._tc.get_or_add_tcPr()
   for old in list(pr.findall(qn('w:shd'))):pr.remove(old)
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F5F5F8' if i%2==0 else 'FFFFFF');pr.append(sh)
staged=work/source.name;doc.save(staged);context['guide_source']=str(source);context['guide_candidate']=str(staged);context['guide_original_sha256']=hashlib.sha256(original).hexdigest();context_state_path.write_text(json.dumps(context,indent=2))
reference=json.loads((root/'design_reference_state.json').read_text());reference.update({'work':str(work),'docx':str(staged)});(root/'design_reference_state.json').write_text(json.dumps(reference,indent=2));print('Staged style-guide version 1.1 with the filter context and partner-only reset behavior.')
